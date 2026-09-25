#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""Телефонный канал замера занятости: лист прозвона -> ответы -> снапшот.

Зачем. У хаусботов (поручение заказчика 07.09.2026: «какие хаусботы лучше всего
сдаются и почему») модуля онлайн-брони нет почти ни у кого: в Самаре — ни у
одного борта, по России — у трёх-четырёх. Единственный дистанционный способ
узнать, сдаётся ли судно, — позвонить под видом гостя и спросить про
конкретные ночи. Этот файл делает такой прозвон ВОСПРОИЗВОДИМЫМ и кладёт
ответы в тот же ряд, что и машинные снимки (SCHEMA (б) ядра, engine "phone",
source_kind "none"), чтобы слой сезонности и сводка читали их наравне с
модулями — с честной пометкой, что это ответ человека, а не календарь.

Две команды:

  sheet  — собрать лист прозвона на неделю:
      phone_channel.py sheet --targets <json> --out-dir <dir> [--date YYYY-MM-DD]
    <json> — houseboat-tracking-targets.json из apply_verdicts.py (ключ
    "selected") ИЛИ произвольный список [{key, unit_name, phones[], region,
    owner, site}]. На выходе xlsx + csv: по строке на судно, телефоны с
    витриной, сценарий вопроса и три окна ночей (ближайшие выходные,
    следующие выходные, одна будняя ночь), в которые звонящий пишет
    свободно / занято / частично / не дозвонились.

  ingest — превратить заполненный лист в fixture-объекты и (по желанию)
    записать снапшот:
      phone_channel.py ingest --sheet <xlsx|csv> --targets <json> --out-dir <dir>
                              [--snapshot-root <root> --recipes <recipes.json>]
    Каждый ответ становится клеткой ночи: свободно -> free, занято -> busy,
    частично -> unknown (звонящий не сказал, какая из двух ночей занята —
    угадывать нельзя), не дозвонились -> объект без клеток, status
    insufficient_data с причиной. Цена ночи, названная по телефону, ложится в
    price клетки. С --snapshot-root вызывается cli.py --fixture … и снапшот
    записывается в указанный каталог (для хаусботов — их собственный корень,
    не боевой глэмпинговый).

Честность: ответ по телефону — «оценка сверху» ещё грубее модуля: «занято»
может значить «не хотим сдавать на одну ночь» или «забыли календарь». Поэтому
reason объекта всегда несёт дату звонка и дословный ответ, а сводка видит
engine "phone" и не смешивает его с модулями.

Правила: звонок делает человек; скрипт ничего не бронирует и никуда не звонит.
"""
from __future__ import annotations

import argparse
import csv
import json
import re
import subprocess
import sys
from datetime import date, datetime, timedelta
from typing import Optional
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:  # pragma: no cover — xlsx необязателен, csv всегда есть
    openpyxl = None

SCRIPTS_DIR = Path(__file__).resolve().parent

ANSWERS = ("свободно", "занято", "частично", "не дозвонились")
STATE_OF = {"свободно": "free", "занято": "busy", "частично": "unknown"}

SCRIPT_TEXT = (
    "Здравствуйте! Хотим снять {unit} с ночёвкой. Подскажите, "
    "{w1} (с пятницы на субботу и с субботы на воскресенье) свободно? "
    "А {w2}? А одну ночь в будни, {w3}? Сколько стоит ночь и есть ли минимальный срок?"
)

# Колонки листа прозвона. Ответные колонки идут после «Сценарий вопроса».
COLS = ["Ключ", "Судно", "Регион", "Собственник / кто продаёт", "Телефоны (витрина)",
        "Сценарий вопроса",
        "Ближайшие выходные: ночи", "Ближайшие выходные: ответ",
        "Следующие выходные: ночи", "Следующие выходные: ответ",
        "Будняя ночь: дата", "Будняя ночь: ответ",
        "Цена ночи, названная (₽)", "Минимальный срок", "Комментарий", "Дата звонка", "Кто звонил"]
WIDTHS = [18, 34, 18, 34, 34, 60, 16, 16, 16, 16, 12, 14, 14, 14, 36, 12, 12]


def slug(text: str) -> str:
    """Ключ цели из названия судна: hb_<транслит>, стабильный и читаемый."""
    table = {"а": "a", "б": "b", "в": "v", "г": "g", "д": "d", "е": "e", "ё": "e", "ж": "zh", "з": "z",
             "и": "i", "й": "y", "к": "k", "л": "l", "м": "m", "н": "n", "о": "o", "п": "p", "р": "r",
             "с": "s", "т": "t", "у": "u", "ф": "f", "х": "h", "ц": "c", "ч": "ch", "ш": "sh", "щ": "sch",
             "ъ": "", "ы": "y", "ь": "", "э": "e", "ю": "yu", "я": "ya"}
    raw = str(text or "")
    head = re.split(r"[—(\[,:]", raw)[0].strip().lower()
    if head.startswith("без имени"):
        # «без имени #2 (AQUASTORIES, FreeDom 75)» — имени нет, зато есть
        # владелец в скобках: ключ строится из него, иначе все безымянные
        # борта страны схлопываются в bez_imeni_1/2/3.
        m = re.search(r"\(([^)]*)\)", raw)
        num = re.search(r"#\s*(\d+)", head)
        head = ((m.group(1) if m else head).split(",")[0].split(":")[0].strip().lower()
                + (" " + num.group(1) if num else ""))
    out = "".join(table.get(ch, ch) for ch in head)
    out = re.sub(r"[^a-z0-9]+", "_", out).strip("_")
    return "hb_" + (out[:40] or "unit")


SCRIPT_TEXT_DAY = (
    "Здравствуйте! Хотим арендовать {unit} на день, без ночёвки. Подскажите, "
    "{w1} (суббота и воскресенье) свободно? А {w2}? А в будний день, {w3}? "
    "Сколько стоит день и есть ли минимальный срок?"
)


def day_windows_for(today: date) -> dict:
    """Режим day (хаусботы: заказчик сдаёт без ночёвки, 07.09.2026): ближайшие
    суббота и воскресенье как ДНИ аренды, следующие сб/вс, один будний день."""
    days_to_sat = (5 - today.weekday()) % 7
    if days_to_sat == 0 or today.weekday() >= 5:
        days_to_sat = (5 - today.weekday()) % 7 or 7
    sat = today + timedelta(days=days_to_sat)
    if sat <= today:
        sat += timedelta(days=7)
    w1 = [sat, sat + timedelta(days=1)]
    w2 = [d + timedelta(days=7) for d in w1]
    wed = sat + timedelta(days=4)  # среда после ближайших выходных
    return {"w1": w1, "w2": w2, "w3": [wed]}


def windows_for(today: date) -> dict:
    """Три окна вопросов: ночи ближайших и следующих выходных, одна будняя ночь.

    Ночь называется датой заезда. Если сегодня уже пятница-воскресенье,
    «ближайшие выходные» — следующая пятница: спрашивать про ночь, которая
    уже идёт, бессмысленно.
    """
    days_to_fri = (4 - today.weekday()) % 7
    if days_to_fri == 0 or today.weekday() > 4:
        days_to_fri = (4 - today.weekday()) % 7 or 7
    fri = today + timedelta(days=days_to_fri)
    if fri <= today:
        fri += timedelta(days=7)
    w1 = [fri, fri + timedelta(days=1)]
    w2 = [d + timedelta(days=7) for d in w1]
    tue = fri + timedelta(days=4)  # вторник после ближайших выходных
    return {"w1": w1, "w2": w2, "w3": [tue]}


def _fmt(d: date) -> str:
    return d.strftime("%d.%m")


def load_targets(path) -> list:
    data = json.load(open(path, encoding="utf-8"))
    rows = data.get("selected") if isinstance(data, dict) else data
    out = []
    for r in rows or []:
        name = r.get("unit_name") or r.get("title") or r.get("username") or ""
        key = r.get("key") or r.get("username") or slug(name)
        phones = r.get("phones") or []
        if not phones and r.get("windows"):
            phones = sorted({(w.get("phone") or "").strip() for w in r["windows"] if (w.get("phone") or "").strip()})
        out.append({"key": key, "unit_name": name, "phones": phones, "region": r.get("region") or "",
                    "owner": r.get("owner") or "", "site": r.get("site") or "",
                    "verdict": r.get("verdict") or ""})
    # ключи обязаны быть уникальными: снапшот пишется по ключу
    seen = {}
    for r in out:
        n = seen.get(r["key"], 0)
        seen[r["key"]] = n + 1
        if n:
            r["key"] = f"{r['key']}_{n + 1}"
    return out


def build_sheet(targets: list, out_dir: Path, today: date, mode: str = "night") -> tuple:
    if mode == "day":
        w = day_windows_for(today)
        w1 = f"{_fmt(w['w1'][0])}–{_fmt(w['w1'][1])}"
        w2 = f"{_fmt(w['w2'][0])}–{_fmt(w['w2'][1])}"
        w3 = f"{_fmt(w['w3'][0])} (среда)"
        script = SCRIPT_TEXT_DAY
    else:
        w = windows_for(today)
        w1 = f"{_fmt(w['w1'][0])}–{_fmt(w['w1'][1] + timedelta(days=1))}"
        w2 = f"{_fmt(w['w2'][0])}–{_fmt(w['w2'][1] + timedelta(days=1))}"
        w3 = f"{_fmt(w['w3'][0])} (вт→ср)"
        script = SCRIPT_TEXT
    rows = []
    for t in targets:
        phones = "\n".join(t["phones"]) if t["phones"] else "телефона на витринах нет — искать"
        unit_short = re.split(r"[—(]", t["unit_name"])[0].strip()
        rows.append([t["key"], t["unit_name"], t["region"], str(t["owner"])[:300], phones,
                     script.format(unit=unit_short, w1=w1, w2=w2, w3=w3),
                     f"{w['w1'][0].isoformat()}, {w['w1'][1].isoformat()}", "",
                     f"{w['w2'][0].isoformat()}, {w['w2'][1].isoformat()}", "",
                     w["w3"][0].isoformat(), "", "", "", "", "", ""])
    out_dir.mkdir(parents=True, exist_ok=True)
    stem = out_dir / f"{today.isoformat()}-лист-прозвона"
    with open(stem.with_suffix(".csv"), "w", encoding="utf-8", newline="") as f:
        wr = csv.writer(f)
        wr.writerow(COLS)
        wr.writerows(rows)
    xlsx = None
    if openpyxl is not None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Прозвон"
        ws.append(COLS)
        for r in rows:
            ws.append(r)
        for i, wdt in enumerate(WIDTHS, 1):
            ws.column_dimensions[get_column_letter(i)].width = wdt
            c = ws.cell(row=1, column=i)
            c.font = Font(bold=True)
            c.fill = PatternFill("solid", fgColor="DDDDDD")
            c.alignment = Alignment(wrap_text=True, vertical="center")
        for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
            for c in row:
                c.alignment = Alignment(wrap_text=True, vertical="top")
        # выпадающие списки ответов
        from openpyxl.worksheet.datavalidation import DataValidation
        dv = DataValidation(type="list", formula1='"' + ",".join(ANSWERS) + '"', allow_blank=True)
        ws.add_data_validation(dv)
        for col in (8, 10, 12):
            for r in range(2, len(rows) + 2):
                dv.add(ws.cell(row=r, column=col))
        ws.freeze_panes = "C2"
        g = wb.create_sheet("Как заполнять")
        g.column_dimensions["A"].width = 110
        for i, t in enumerate([
            f"Лист прозвона хаусботов на неделю с {today.isoformat()} ({'ДНИ дневной аренды, без ночёвки' if mode == 'day' else 'ночи'}). Звоним под видом гостя, ничего не бронируем.",
            "",
            "В колонках «ответ» выбираем из списка: свободно / занято / частично / не дозвонились. «Частично» — когда из двух ночей выходных свободна одна "
            "(в ряд ложится «неизвестно», потому что мы не знаем, какая именно). Цену пишем числом в рублях за ночь, как назвали.",
            "",
            "Заполненный файл возвращается скриптом phone_channel.py ingest в ряд занятости: каждая ночь становится клеткой с датой звонка и дословным ответом, "
            "поэтому в «Комментарий» стоит записать всё, что сказали: «только от двух суток», «ночью не сдаём», «до ноября всё занято» — это и есть материал для «почему».",
            "",
            ("Окна недели: ближайшие сб и вс как дни аренды, следующие сб и вс, одна среда. Даты в колонках «ночи» — это ДНИ аренды." if mode == "day" else
             "Окна недели: ближайшие выходные (ночи пт→сб и сб→вс), следующие выходные, одна будняя ночь (вт→ср). Даты в колонках «ночи» — даты заезда."),
        ], 1):
            c = g.cell(row=i, column=1, value=t)
            c.alignment = Alignment(wrap_text=True, vertical="top")
        xlsx = stem.with_suffix(".xlsx")
        wb.save(xlsx)
    return stem.with_suffix(".csv"), xlsx, len(rows)


def read_sheet(path) -> list:
    path = Path(path)
    if path.suffix.lower() == ".xlsx":
        if openpyxl is None:
            sys.exit("openpyxl не установлен, а лист — xlsx")
        wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        ws = wb[wb.sheetnames[0]]
        it = ws.iter_rows(values_only=True)
        header = [str(h or "").strip() for h in next(it)]
        rows = [dict(zip(header, [("" if v is None else v) for v in r])) for r in it if r and r[0]]
    else:
        with open(path, encoding="utf-8", newline="") as f:
            rows = list(csv.DictReader(f))
    return rows


def _nights(text, bad: Optional[list] = None) -> list:
    """Ночи из ячейки листа -> ['YYYY-MM-DD', ...].

    Лист заполняет человек в Excel (ревью 14.09.2026): ячейку он превращает в
    дату (openpyxl отдаёт datetime, а str() даёт '2026-09-22 00:00:00'), руками
    пишут «22.09.2026». Раньше строка уходила ключом клетки как есть, и один такой
    объект ронял запись всего снапшота. Неразобранное — в bad, в клетки не идёт.
    """
    if isinstance(text, datetime):
        return [text.date().isoformat()]
    if isinstance(text, date):
        return [text.isoformat()]
    out = []
    for piece in str(text or "").split(","):
        piece = piece.strip()
        if not piece:
            continue
        m = re.fullmatch(r"(\d{4})-(\d{2})-(\d{2})(?:[ T].*)?", piece)
        if m:
            out.append(f"{m.group(1)}-{m.group(2)}-{m.group(3)}")
            continue
        m = re.fullmatch(r"(\d{1,2})\.(\d{1,2})\.(\d{4})", piece)
        if m:
            out.append(f"{m.group(3)}-{int(m.group(2)):02d}-{int(m.group(1)):02d}")
            continue
        if bad is not None:
            bad.append(piece)
    return out


def _price(text):
    """Первая сумма из ячейки «Цена ночи, названная (₽)».

    Колонка — свободный текст (ревью 14.09.2026): прежний разбор склеивал все цифры,
    и «25 000 – 30 000» становилось 2 500 030 000 ₽, «4 500 ₽, мин. 2 ночи» — 45 002.
    Берём первое число с разделителями тысяч; дробную часть отбрасываем."""
    m = re.search(r"\d{1,3}(?:[ \u00a0\u202f]\d{3})+|\d+", str(text or ""))
    if not m:
        return None
    return int(re.sub(r"\D", "", m.group(0)))


def fixtures_from_sheet(rows: list, targets: list, called_on: str) -> list:
    """Строки листа -> fixture-объекты SCHEMA (б)."""
    by_key = {t["key"]: t for t in targets}
    out = []
    for r in rows:
        key = str(r.get("Ключ") or "").strip()
        if not key:
            continue
        t = by_key.get(key, {"unit_name": r.get("Судно"), "site": ""})
        cells = {}
        answers = []
        unparsed = []
        price = _price(r.get("Цена ночи, названная (₽)"))
        for nights_col, ans_col, label in (("Ближайшие выходные: ночи", "Ближайшие выходные: ответ", "ближайшие выходные"),
                                           ("Следующие выходные: ночи", "Следующие выходные: ответ", "следующие выходные"),
                                           ("Будняя ночь: дата", "Будняя ночь: ответ", "будняя ночь")):
            ans = str(r.get(ans_col) or "").strip().lower()
            if not ans:
                continue
            answers.append(f"{label}: {ans}")
            state = STATE_OF.get(ans)
            if state is None:
                continue  # «не дозвонились» и всё незнакомое клеток не даёт
            for night in _nights(r.get(nights_col), bad=unparsed):
                # Судно — один физический юнит: пара фонда честна и говорит
                # сводке «домик», а не «тип с неснятым фондом».
                cell = {"state": state}
                if state != "unknown":
                    cell["units_total"] = 1
                    cell["units_free"] = 1 if state == "free" else 0
                if price and state == "free":
                    cell["price"] = price
                cells[night] = cell
        when = str(r.get("Дата звонка") or called_on).strip() or called_on
        comment = str(r.get("Комментарий") or "").strip()
        who = str(r.get("Кто звонил") or "").strip()
        reason_bits = [f"прозвон под видом гостя {when}" + (f" ({who})" if who else "")]
        if answers:
            reason_bits.append("ответы: " + "; ".join(answers))
        if r.get("Минимальный срок"):
            reason_bits.append(f"минимальный срок: {r.get('Минимальный срок')}")
        if comment:
            reason_bits.append("сказали: " + comment)
        if unparsed:
            reason_bits.append("не разобраны даты в листе: " + ", ".join(unparsed))
        obj = {"username": key, "site": t.get("site") or "", "engine": "phone", "source_kind": "none",
               "granularity": "aggregate", "units": {"__aggregate__": cells},
               "source_urls": [], "reason": ". ".join(reason_bits),
               "status": "ok" if cells else "insufficient_data"}
        if not cells:
            obj["reason"] = obj["reason"] + ". Клеток нет: не дозвонились или ответ не разобран"
        out.append(obj)
    return out


def main(argv=None) -> int:
    ap = argparse.ArgumentParser(description=__doc__.split("\n")[0])
    sub = ap.add_subparsers(dest="cmd", required=True)
    s = sub.add_parser("sheet")
    s.add_argument("--targets", required=True)
    s.add_argument("--out-dir", required=True)
    s.add_argument("--date", default=None)
    s.add_argument("--mode", choices=("night", "day"), default="night",
                   help="day — дни дневной аренды (хаусботы без ночёвки), night — ночи")
    i = sub.add_parser("ingest")
    i.add_argument("--sheet", required=True)
    i.add_argument("--targets", required=True)
    i.add_argument("--out-dir", required=True, help="куда положить fixture-файлы")
    i.add_argument("--date", default=None, help="дата звонка по умолчанию")
    i.add_argument("--snapshot-root", default=None, help="записать снапшот через cli.py --fixture")
    i.add_argument("--recipes", default=None)
    a = ap.parse_args(argv)
    today = date.fromisoformat(a.date) if a.date else date.today()
    targets = load_targets(a.targets)
    if a.cmd == "sheet":
        csv_path, xlsx, n = build_sheet(targets, Path(a.out_dir), today, mode=a.mode)
        print(f"лист прозвона: судов {n}; csv {csv_path}" + (f"; xlsx {xlsx}" if xlsx else ""))
        return 0
    rows = read_sheet(a.sheet)
    fixtures = fixtures_from_sheet(rows, targets, today.isoformat())
    out_dir = Path(a.out_dir)
    out_dir.mkdir(parents=True, exist_ok=True)
    paths = []
    for obj in fixtures:
        p = out_dir / f"{obj['username']}.json"
        json.dump(obj, open(p, "w", encoding="utf-8"), ensure_ascii=False, indent=1)
        paths.append(p)
    with_cells = sum(1 for o in fixtures if o["status"] == "ok")
    print(f"fixture-объектов {len(fixtures)} (с клетками {with_cells}) -> {out_dir}")
    if a.snapshot_root and paths:
        cmd = [sys.executable, str(SCRIPTS_DIR / "cli.py"), "--snapshot-root", a.snapshot_root]
        if a.recipes:
            cmd += ["--recipes", a.recipes]
        for p in paths:
            cmd += ["--fixture", str(p)]
        print("снапшот:", " ".join(cmd[:4]), f"... ({len(paths)} fixture)")
        rc = subprocess.call(cmd)
        return rc
    return 0


if __name__ == "__main__":
    sys.exit(main())
