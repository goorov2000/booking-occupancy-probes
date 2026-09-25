# -*- coding: utf-8 -*-
"""Поюнитный отчёт и отчёт о сезонности (тикеты 14 и 15).

Зачем он есть
-------------
Заказчик просил это трижды — 26.08, 27.08 и 28.08: «мы не видим, какие домики
сдаются, какие нет; месячные проценты слишком размыты». Данные для ответа лежат
в снапшотах с 14.08 (сетка «юнит × ночь» сохраняется целиком), а сводка
`/glamping-occupancy` показывает только месячный процент ПО ОБЪЕКТУ — ровно ту
размытость, на которую жалоба. Не хватало не данных, а рендера. Здесь он.

Что этот модуль обязан не соврать
---------------------------------
1. **Три рода строк, три подписи, никогда не в одном столбце** (решение Р3
   спеки). «Домик: <имя>» — движок адресует физический номер. «Тип: <имя> —
   продано N из M» — движок торгует категорией, фонд снят. «Тип: <имя> — фонд
   не снят» — сколько домиков за именем, неизвестно.
   Род определяется ДВИЖКОМ и наличием фонда, а НЕ именем юнита: у `wood_glamp`
   юниты называются «Сфера 1»…«Сфера 5» и выглядят домиками, но фонд не снят, а
   у `a_ureki` «Белый Дом» — домик, зато соседний «Этнодом» — четыре домика под
   одним именем. Строка второго рода не называется домиком даже при M=1.
   Род берётся из слоя сезонности (`ledger`), а не считается заново: двух
   источников одного признака в этом инструменте уже хватило — из-за трёх
   копий предиката «готова ли тема» на заводе 10 утверждённых рецептов молча не
   давали сценариев.
2. **«Продано вперёд» и «фактически занято» — две разные цифры и НИКОГДА не
   одна колонка.** Первая — глубина продаж из последнего снимка (что стоит в
   календаре на будущие ночи). Вторая — реализованная занятость прошедших
   ночей из ряда (`ledger`, правило Р4: состояние ночи из последнего снимка,
   где она ещё была в горизонте). Смешать их — значит выдать план за факт.
3. **Ряд короче месяца фактом не печатается.** Слежка идёт с 14.08, поэтому
   август у всех — восемнадцать ночей из тридцати одной; такая цифра выходит с
   подписью «неполный ряд», а не как месячный факт (приёмка тикета 15).
4. **Месяц, где фонд снят не на все ночи, помечается.** С тикета 06 фонд типов
   спрашивается на ближние 45 ночей — дальше тип из трёх домиков весит один, и
   процент занижен методически, а не потому что не бронируют (приёмка тикета 14).
   Помечается ВЕЗДЕ, где этот процент напечатан: в поюнитной ячейке, на листе
   «Сезонность по месяцам», в строке медианы и в markdown-выжимке. Пометка,
   дошедшая только до одной из четырёх таблиц, — это её отсутствие: заказчик
   читает лист сезонности, а не поюнитную ячейку.
5. **Ничто не исчезает.** Юнит, у которого сегодня нет данных или которого нет
   в реестре целей, остаётся честной строкой с причиной (правило №2 завода).
6. **Дыра в слое сезонности называется вслух.** Прогон, не попавший в ряд, и
   индекс, обещающий больше строк, чем лежит в файле, — это дыра в ЦИФРАХ
   отчёта, а не служебная мелочь: ряд без пропущенных суток выглядит как
   честный. Формулировку даёт сам слой (`ledger.hole_note`), отчёт печатает её
   оператору в вывод прогона, а заказчику — на лист «Чего пока нет» и в
   «Как читать».

Разрез по региону
-----------------
Поле `region` в `targets.json` ставит владелец отчётов при сборке самарской
партии. Пока его нет, разрез недоступен — и это говорится строкой, а не падает
(`ledger.region_coverage`). Появится — самарские и подмосковные объекты
разъедутся по разным листам и в разные медианы: складывать их в одну нельзя.

Формат доставки
---------------
Тот же, что у отчёта по хаусботам, который заказчик принял 04.09: xlsx (openpyxl),
первым листом «Как читать» человеческим текстом, у каждой таблицы пустые
колонки «Вердикт заказчика» и «Комментарий», отдельные листы на регион. Рядом —
версионируемая выжимка в markdown. Существующая `occupancy.md` не трогается:
она решает другую задачу («снимок перед интервью» по 21 кандидату) и уже 37 КБ.

Сети здесь нет и быть не может: модуль читает только диск.
"""
from __future__ import annotations

import argparse
import calendar
import sys
import zlib
from datetime import date, datetime
from pathlib import Path
from typing import Optional

import build_summary as bs
import ledger
import occupancy_core as core

# --- роды строк (решение Р3). Имена — те же, что в слое сезонности ---------
KIND_HOME = ledger.BASIS_HOME
KIND_TYPE = ledger.BASIS_TYPE
KIND_TYPE_NO_CAPACITY = ledger.BASIS_TYPE_NO_CAPACITY
KINDS = (KIND_HOME, KIND_TYPE, KIND_TYPE_NO_CAPACITY)

# Как род называется в человеческом тексте отчёта.
KIND_TITLES = {
    KIND_HOME: "домик — движок адресует физический номер",
    KIND_TYPE: "тип с известным фондом — движок отвечает «продано N из M»",
    KIND_TYPE_NO_CAPACITY: "тип, у которого фонд не снят",
}

# --- шапка таблицы --------------------------------------------------------
COL_NUM = "№"
COL_OBJECT = "Объект"
COL_SITE = "Сайт"
COL_KIND = "Род строки"
COL_ENGINE = "Модуль брони"
COL_CAPACITY = "Домиков в юните"
COL_NIGHTS = "Ночей в ряду"
COL_CHECKED = "Снято (дата)"
COL_STATE = "Состояние съёма"
COL_NOTES = "Оговорки"
COL_VERDICT = "Вердикт заказчика"
COL_COMMENT = "Комментарий"

FACT_PREFIX = "Факт: "
AHEAD_PREFIX = "Продано вперёд: "

ALL_TARGETS_SHEET = "Юниты (все цели)"
NO_REGION_SHEET = "Юниты — вне реестра"
HOW_SHEET = "📖 Как читать"
SEASON_SHEET = "Сезонность по месяцам"
PACE_SHEET = "Темп бронирования"
LIMITS_SHEET = "Чего пока нет"
NODATA_SHEET = "Объекты без цифр"

# Сколько месяцев факта показываем. Больше года в одну таблицу не влезает по
# ширине, а меньше полугода прячет сезонность — ради которой всё и затевалось.
DEFAULT_FACT_MONTHS = 6

DEFAULT_OUT_DIR = core.OCCUPANCY_ROOT / "out"
DEFAULT_MD_DIR = (core.REPO_ROOT / "docs" / "research"
                  / "2026-08-10-glamping-market")


# ---------------------------------------------------------------------------
# Род строки и её подпись
# ---------------------------------------------------------------------------

def row_kind(engine: str, granularity: str, unit: str,
             capacity: Optional[int]) -> str:
    """Род строки по движку и фонду (решение Р3), НИКОГДА не по имени юнита.

    Своей копии правила здесь нет нарочно: род обязан быть один на всю
    систему, иначе отчёт и ряд подпишут одну и ту же строку по-разному.
    Публичного имени у предиката в `ledger` пока нет — попросить владельца
    слоя (см. notes_for_others), а до тех пор берём его же функцию.
    """
    return ledger._row_basis(engine or "", granularity or "", unit, capacity)


def row_label(kind: str, unit: str, *, capacity: Optional[int] = None,
              sold: Optional[int] = None, night: Optional[str] = None,
              no_fund: bool = False) -> str:
    """Подпись строки — одна из трёх, и по ней сразу видно род.

    У типа с фондом в подпись идёт «продано N из M» на ближайшую СНЯТУЮ ночь,
    и ОБА числа берутся из одной клетки: это ровно то, что ответил движок.
    Склеивать «продано» из клетки без снятого фонда (там любой тип весит один
    домик) с фондом из максимума за всю историю нельзя — получится пара чисел
    из двух источников, поданная как один ответ.
    Ближайшая ночь не снята (окно продаж закрыто, отказ хоста) либо фонд на
    ближние ночи не спрашивался — подпись остаётся типовой и называет только
    фонд: домиком строка не становится никогда.
    """
    if kind == KIND_HOME:
        return f"Домик: {unit}"
    if kind == KIND_TYPE_NO_CAPACITY:
        return f"Тип: {unit} — фонд не снят"
    total = capacity if capacity is not None else "?"
    homes = (f"{total} {_homes_word(total)}" if isinstance(total, int)
             else f"{total} домиков")
    if sold is None and no_fund:
        return f"Тип: {unit} — {homes}, фонд на ближайшие ночи не снят"
    if sold is None:
        return f"Тип: {unit} — {homes}, ближайшая ночь не снята"
    tail = f" на ночь {_fmt_day(night)}" if night else ""
    return f"Тип: {unit} — продано {sold} из {total}{tail}"


# ---------------------------------------------------------------------------
# Мелкая типографика
# ---------------------------------------------------------------------------

def _nights_word(n: int) -> str:
    """«1 ночь», «2 ночи», «5 ночей» — иначе оговорка читается как машинная."""
    if n % 10 == 1 and n % 100 != 11:
        return "ночь"
    if n % 10 in (2, 3, 4) and n % 100 not in (12, 13, 14):
        return "ночи"
    return "ночей"


def _nights_dative(n: int) -> str:
    """Дательный для «по N ноч…»: «по 1 ночи», «по 17 ночам»."""
    return "ночи" if n % 10 == 1 and n % 100 != 11 else "ночам"


def _homes_word(n: int) -> str:
    """«1 домик», «3 домика», «7 домиков» — подпись читает человек."""
    if n % 10 == 1 and n % 100 != 11:
        return "домик"
    if n % 10 in (2, 3, 4) and n % 100 not in (12, 13, 14):
        return "домика"
    return "домиков"


def _agree(n: int, one: str, many: str) -> str:
    """Согласование глагола с числом: «1 ночь снята», «5 ночей сняты»."""
    return one if n % 10 == 1 and n % 100 != 11 else many


def _fmt_pct(pct) -> str:
    return "—" if pct is None else f"{pct:.1f}%"


def _fmt_day(value) -> str:
    """ISO-дата или ISO-время -> «18.08.2026»; мусор -> пустая строка."""
    if not value:
        return ""
    try:
        d = datetime.fromisoformat(str(value)).date()
    except ValueError:
        try:
            d = date.fromisoformat(str(value)[:10])
        except ValueError:
            return ""
    return d.strftime("%d.%m.%Y")


def _month_label(month: str) -> str:
    year, mon = (int(x) for x in month.split("-"))
    return f"{core.MONTH_LABELS_RU[mon]} {year}"


def _days_in_month(month: str) -> int:
    year, mon = (int(x) for x in month.split("-"))
    return calendar.monthrange(year, mon)[1]


# Ячейке таблицы хватает «сен 2026» (`core.MONTH_LABELS_RU`), а человеческому
# тексту — нет: «первый полный месяц — сен 2026» читается как машинный обрывок.
MONTHS_NOM = {1: "январь", 2: "февраль", 3: "март", 4: "апрель", 5: "май",
              6: "июнь", 7: "июль", 8: "август", 9: "сентябрь",
              10: "октябрь", 11: "ноябрь", 12: "декабрь"}

MONTHS_GEN = {1: "января", 2: "февраля", 3: "марта", 4: "апреля", 5: "мая",
              6: "июня", 7: "июля", 8: "августа", 9: "сентября",
              10: "октября", 11: "ноября", 12: "декабря"}


def _month_words(month: str) -> str:
    """«2026-09» -> «сентябрь 2026» — для текста, который читает человек."""
    year, mon = (int(x) for x in month.split("-"))
    return f"{MONTHS_NOM[mon]} {year}"


def _month_shift(month: str, delta: int) -> str:
    """«2026-08» + 2 -> «2026-10». Сроки считаются, а не дописываются руками."""
    year, mon = (int(x) for x in month.split("-"))
    total = year * 12 + (mon - 1) + delta
    return f"{total // 12:04d}-{total % 12 + 1:02d}"


def _ready_day(month: str) -> str:
    """Когда месяц закрыт: первое число СЛЕДУЮЩЕГО месяца, словами."""
    year, mon = (int(x) for x in _month_shift(month, 1).split("-"))
    return f"1 {MONTHS_GEN[mon]} {year} года"


def _first_full_month(first_night: Optional[str]) -> Optional[str]:
    """Первый месяц, который слежка застала целиком (а не с середины)."""
    if not first_night:
        return None
    try:
        day = date.fromisoformat(str(first_night)[:10])
    except ValueError:
        return None
    month = day.strftime("%Y-%m")
    return month if day.day == 1 else _month_shift(month, 1)


def md_cell(text) -> str:
    """Чужой текст в ячейку markdown: без переводов строк и без сырых `|`.

    Имя юнита приходит от чужого движка. Одна вертикальная черта в нём —
    и таблица отчёта разъезжается целиком; ровно этот дефект уже ловили в
    сводке (ревью волны 1), незачем повторять его здесь.
    """
    if text is None:
        return ""
    return (str(text).replace("\r", " ").replace("\n", " ")
            .replace("|", "\\|").strip())


# ---------------------------------------------------------------------------
# Ячейки: факт, продажи вперёд, темп
# ---------------------------------------------------------------------------

def fact_cell(fact: dict) -> str:
    """Ячейка фактической занятости месяца.

    Ряд короче месяца НЕ печатается как месячный факт (приёмка тикета 15):
    цифра выходит с подписью «неполный ряд: … по N ночам из M», и спутать её
    с месячным итогом нельзя.

    Второе, что обязана сказать ячейка, — что её знаменатель СМЕШАН из двух
    оснований: часть ночей движок отдал с фондом (тип весит три домика), часть
    без него (тот же тип весит один). Один процент по такому знаменателю
    несравним с соседней строкой, и разница между ними будет методической, а
    не про брони. Ровно эту ошибку тикет 08 чинил на уровне объекта.
    """
    if not fact:
        return ""
    cuts = fact["cuts"]["all"]
    if cuts["known"] == 0:
        if fact.get("quota_excluded"):
            return "квота агрегатора — это доля площадки, а не загрузка объекта"
        if fact.get("sales_not_open"):
            return "продажи не открыты"
        return "нет данных"
    body = (f"{_fmt_pct(cuts['pct'])} "
            f"({cuts['busy']} из {cuts['known']} домико-ночей)")
    if not fact.get("complete"):
        # «Посчитано по N ночам» — это ночи, вошедшие в ПРОЦЕНТ, а не длина
        # ряда: `nights` считает и unknown, и «продажи не открыты», которых в
        # знаменателе нет, и завышал надёжность цифры (17 ночей вместо 12).
        counted = fact.get("measured_nights") or fact["nights"]
        body = (f"неполный ряд: {body}, посчитано по {counted} "
                f"{_nights_dative(counted)} из {fact['days_in_month']}")
    short = fact.get("inventory_short") or 0
    if short:
        cap = fact.get("capacity")
        weight = (f" своих {cap} {_homes_word(cap)}" if cap else " всего типа")
        body += (f"; фонд снят не на все ночи ({short} из "
                 f"{fact['measured_nights']} снятых): в них юнит весит меньше"
                 f"{weight} — процент занижен")
    return body


def ahead_cell(ahead: dict, gap: bool = False) -> str:
    """Ячейка глубины продаж вперёд. Это НЕ факт и стоит в другой колонке.

    По скольким ночам посчитано — часть самой цифры, а не сноска: «октябрь
    0.0%» по ОДНОЙ контрольной ночи и «октябрь 0.0%» по тридцати одной — это
    два разных утверждения. Колонка факта так и написана с самого начала;
    здесь то же самое.
    """
    if not ahead:
        return ""
    cuts = ahead["cuts"]["all"]
    if cuts["known"] == 0:
        if ahead.get("sales_not_open"):
            return (f"продажи не открыты ({ahead['sales_not_open']} "
                    f"{_nights_word(ahead['sales_not_open'])})")
        return "нет данных"
    body = (f"{_fmt_pct(cuts['pct'])} "
            f"({cuts['busy']} из {cuts['known']} домико-ночей)")
    nights, days = ahead.get("nights"), ahead.get("days_in_month")
    if nights and days and nights < days:
        body += (f", посчитано по {nights} {_nights_dative(nights)} "
                 f"из {days}")
    if gap:
        body += " — фонд снят не на все ночи, процент занижен"
    return body


def pace_cell(point: dict) -> str:
    """Ячейка точки кривой темпа: цифра либо честная ПРИЧИНА её отсутствия.

    Причин у пустой точки четыре, и они разные. «Мало ночей» — точка стоит на
    меньшем наборе ночей, чем соседние (у глубоких lead истории просто нет).
    «Ночи не сняты» — набор полный, но ни одной измеренной клетки в нём нет
    (так живёт `ok_reka`: 21 снимок подряд без единой проверенной ночи).
    «Продажи не открыты» — модуль закрыл окно. «Фонд не снят» — точку съел
    ближний горизонт фонда (тикет 06). Валить их в одну подпись нельзя:
    «мало ночей: 17 из 17» — это ложь про полный набор.

    Точка С цифрой тоже несёт своё основание: на живом слое 180 из 242
    напечатанных процентов стояли на неполном наборе ночей и читались как
    сравнимые с соседними. Порог `ledger.DEFAULT_MIN_COVERAGE` отсекает только
    совсем тонкие точки, поэтому у остальных набор ночей пишется в скобках
    прямо рядом с цифрой — иначе падение кривой слева направо читается как
    темп продаж, хотя часть его — разная глубина истории.
    """
    if not point:
        return ""
    if point.get("pct") is not None:
        marks = []
        if point.get("partial"):
            marks.append(f"по {point['nights']} "
                         f"{_nights_dative(point['nights'])} из "
                         f"{point['nights_base']}")
        if point.get("inventory_missing"):
            marks.append(f"без снятого фонда {point['inventory_missing']} "
                         f"юнит-{_nights_word(point['inventory_missing'])}")
        if marks:
            return f"{_fmt_pct(point['pct'])} ({'; '.join(marks)})"
        return _fmt_pct(point["pct"])
    if point.get("known"):
        return f"мало ночей: {point['nights']} из {point['nights_base']}"
    if point.get("sales_not_open"):
        return "продажи не открыты"
    if point.get("unknown"):
        return "ночи не сняты"
    if point.get("inventory_missing"):
        return "фонд не снят"
    return "—"


# ---------------------------------------------------------------------------
# Сборка данных
# ---------------------------------------------------------------------------

def _latest_nights(rows: list[dict]) -> list[dict]:
    """По одной строке на «ночь+объект+юнит» — последнюю (журнал дописывается)."""
    seen: dict[tuple, dict] = {}
    for row in rows:
        seen[(row["night"], row["username"], row["unit"])] = row
    return list(seen.values())


def _kept_nights(rows: list[dict], month: Optional[str] = None) -> list[dict]:
    """Строки ряда, которые РЕАЛЬНО попали в факт (то же сито, что у `ledger`).

    Своей копии правила отбора здесь нет нарочно: ночи с опозданием и квоты
    агрегаторов выбрасывает `ledger`, и если отчёт начнёт решать это заново,
    его знаменатель разъедется с напечатанным процентом. Публичного имени у
    сита пока нет — попросить владельца слоя (см. notes_for_others).
    """
    kept, _dropped = ledger._select(rows, None, None, month, None, False,
                                    False)
    return kept


def _unit_ahead(cells: dict, months: list[str]) -> dict:
    """Глубина продаж вперёд по месяцам — считает ЯДРО, а не этот модуль."""
    out = {}
    for m in core.aggregate({"unit": cells}, months):
        nights = sum(1 for day, cell in cells.items()
                     if day.startswith(m["month"])
                     and cell.get("state") in ("free", "busy"))
        out[m["month"]] = {"month": m["month"], "label": m["label"],
                           "cuts": m["cuts"], "unknown": m["unknown"],
                           "sales_not_open": m["sales_not_open"],
                           "nights": nights,
                           "days_in_month": _days_in_month(m["month"])}
    return out


def _unit_fact(nights: list[dict], months: list[str],
               capacity: Optional[int] = None) -> dict:
    """Реализованная занятость юнита по месяцам — из ряда (`ledger`).

    Месяц, у которого ночи в ряду ЕСТЬ, но все они выброшены (квота
    агрегатора, съём с опозданием), всё равно попадает в ответ: пустая
    ячейка без причины читается как «объект простаивал», а это не так.

    Вместе с процентом считается `inventory_short` — сколько СНЯТЫХ ночей
    месяца пришли с фондом меньше фонда юнита. Это и есть смешанный
    знаменатель: у таких ночей тип из трёх домиков весит один.
    """
    have = {row["night"][:7] for row in nights}
    out = {}
    for month in months:
        if month not in have:
            continue
        res = ledger.realized(month=month, rows=nights)
        days = _days_in_month(month)
        measured = [r for r in _kept_nights(nights, month)
                    if r.get("state") in ("free", "busy")]
        short = [r for r in measured
                 if capacity and (r.get("units_total") or 0) < capacity]
        res.update(month=month, label=_month_label(month),
                   days_in_month=days, complete=res["nights"] >= days,
                   capacity=capacity,
                   measured_nights=len({r["night"] for r in measured}),
                   inventory_short=len({r["night"] for r in short}))
        out[month] = res
    return out


def _row_notes(row: dict, months_ahead: list[str]) -> str:
    """Оговорки строки одной фразой: почему цифре нельзя верить как есть."""
    bits = []
    if not row["in_grid"]:
        bits.append("юнита нет в сегодняшней сетке")
    if row["quota"]:
        bits.append("снято квотой площадки-агрегатора: это доля площадки, а "
                    "не загрузка объекта, и в факт она не идёт")
    # Месяц с неполным фондом называется один раз — и будущий (горизонт
    # фонда, тикет 06), и прошедший (ночи ряда, снятые без фонда): читателю
    # важен не источник дыры, а то, что процент месяца занижен методически.
    gaps = {m for m in months_ahead if row["inventory_gap"].get(m)}
    gaps |= {m for m, f in row["fact"].items() if f.get("inventory_short")}
    if gaps:
        bits.append("фонд снят не на все ночи: "
                    + ", ".join(_month_label(m) for m in sorted(gaps))
                    + " — процент этих месяцев занижен методически")
    closed = sum(a.get("sales_not_open", 0) for a in row["ahead"].values())
    if closed:
        bits.append(f"{closed} {_nights_word(closed)} вперёд "
                    f"{_agree(closed, 'закрыта', 'закрыты')} "
                    f"(продажи не открыты) — это не занятость")
    dropped = sum(f.get("gap_excluded", 0) for f in row["fact"].values())
    if dropped:
        bits.append(f"{dropped} {_nights_word(dropped)} ряда "
                    f"{_agree(dropped, 'снята', 'сняты')} с опозданием и в "
                    f"факт {_agree(dropped, 'не пошла', 'не пошли')}")
    if row["kind"] == KIND_TYPE_NO_CAPACITY:
        bits.append("фонд типа не снят: тип занят, только когда продан "
                    "последний домик — занижение")
    return "; ".join(bits)


def collect(*, snapshot_root=None, targets=None, recipes=None, ledger_dir=None,
            today_: Optional[date] = None, months_ahead: int = 3,
            months_fact: int = DEFAULT_FACT_MONTHS) -> dict:
    """Собрать модель отчёта. Ничего не пишет на диск — только читает.

    `targets` и `recipes` принимают и путь, и уже прочитанный объект: отчёт
    зовут и из CLI, и из тестов, и незачем ради этого писать файл.
    """
    snapshot_root = Path(snapshot_root or core.DEFAULT_SNAPSHOT_ROOT)
    today = today_ if today_ is not None else core.today()
    targets = _as_targets(targets)
    recipes = _as_recipes(recipes)

    ahead_months = core.summary_months(today, count=months_ahead)
    history = bs.load_history(snapshot_root, depth=2)
    obj_rows = bs.build_rows(targets, recipes, history, ahead_months,
                             today_=today)
    objects = {row["username"]: row for row in obj_rows}

    nights, pace_rows, capacity, index, ledger_error = _read_layer(ledger_dir)
    # Дыра слоя (прогон, не попавший в ряд; индекс, обещающий больше строк,
    # чем в файле) — это дыра в ЦИФРАХ отчёта, и молчать о ней нельзя.
    ledger_hole = ledger.hole_note(index, nights_read=len(nights),
                                   pace_read=len(pace_rows)) if index else ""
    fact_months = _fact_months(nights, months_fact)

    by_user_unit: dict[tuple, list] = {}
    for row in _latest_nights(nights):
        by_user_unit.setdefault((row["username"], row["unit"]), []).append(row)

    regions = {t["username"]: t.get("region") for t in targets
               if isinstance(t, dict) and t.get("username")}
    units = _unit_rows(objects, by_user_unit, capacity, ahead_months,
                       fact_months, regions)

    data = {
        "generated_for": today,
        "snapshot_root": str(snapshot_root),
        "months_ahead": ahead_months,
        "months_fact": fact_months,
        "objects": objects,
        "units": units,
        "targets": targets,
        "by_kind": {kind: sum(1 for u in units if u["kind"] == kind)
                    for kind in KINDS},
        "region": ledger.region_coverage(targets),
        "ledger_error": ledger_error,
        "ledger_hole": ledger_hole,
        "ledger_index": index,
        # Границы ряда: по ним считаются сроки на листе «Чего пока нет» —
        # какой месяц слежка застала целиком и когда он будет закрыт.
        "ledger_span": {"first": min((r["night"] for r in nights),
                                     default=None),
                        "last": max((r["night"] for r in nights),
                                    default=None)},
        "_pace_rows": pace_rows,
    }
    data["no_data"] = _no_data_rows(objects, units)
    data["seasonality"] = _seasonality_rows(by_user_unit, fact_months, units)
    data["market"] = _market_rows(nights, fact_months, targets,
                                  data["region"], data["seasonality"])
    data["pace"] = _pace_curves(pace_rows, units, fact_months + ahead_months)
    return data


def _no_data_rows(objects: dict, units: list[dict]) -> list[dict]:
    """Цели, у которых нет НИ ОДНОГО юнита: честная строка с причиной.

    Правило №2 завода: объект без модуля бронирования (или со сломанным
    рецептом) остаётся в отчёте строкой «нет данных» с причиной, а не
    исчезает. Держим его на отдельном листе, чтобы в поюнитной таблице
    остались только настоящие юниты и её три подписи не размывались
    четвёртой, у которой юнита нет вовсе.
    """
    with_units = {u["username"] for u in units}
    rows = []
    for username, obj in objects.items():
        if username in with_units:
            continue
        rows.append({
            "username": username,
            "site": obj.get("site") or "",
            "engine": obj.get("engine") or "",
            "status": obj.get("status") or "нет снимка",
            "reason": obj.get("reason") or "",
            "source_label": obj.get("source_label") or "",
            "checked_at": obj.get("checked_at"),
        })
    return rows


NODATA_HEADERS = ["Объект", "Сайт", "Модуль брони", "Почему цифр нет",
                  "Состояние съёма", "Последний снимок", COL_VERDICT,
                  COL_COMMENT]


def nodata_sheet(data: dict) -> tuple:
    rows = []
    for r in data["no_data"]:
        reason = r["reason"] or r["status"]
        # Сводка часто пересказывает ту же причину слово в слово; печатать её
        # дважды — значит заставлять читателя сверять две одинаковые простыни.
        state = "" if reason and reason in r["source_label"] \
            else r["source_label"]
        rows.append([r["username"], r["site"], r["engine"] or "не опознан",
                     reason, state,
                     _fmt_day(r["checked_at"]) or "снимка не было", "", ""])
    return NODATA_HEADERS, rows


def _as_targets(targets):
    if targets is None:
        return core.load_targets(core.DEFAULT_TARGETS)
    if isinstance(targets, (str, Path)):
        return core.load_targets(targets)
    return list(targets)


def _as_recipes(recipes):
    if recipes is None:
        return core.load_recipes(core.DEFAULT_RECIPES)
    if isinstance(recipes, (str, Path)):
        return core.load_recipes(recipes)
    return dict(recipes)


def _read_layer(ledger_dir):
    """Ряд, темп и фонд из слоя сезонности. Слой не читается — строка, не падение.

    Отчёт без ряда всё равно полезен (глубина продаж вперёд считается из
    снапшота), и терять его целиком из-за несобранного производного слоя
    нельзя: слой лечится одной командой `ledger.rebuild`.

    Ловится не только «слоя нет» (`LedgerError`), но и повреждение файла:
    оборванный член gzip даёт `EOFError`/`BadGzipFile`, недописанная строка —
    `JSONDecodeError`. Дозапись слоя идёт в хвосте планового прогона, туда же
    прилетает SIGTERM по `TimeoutStartSec`, — то есть один неудачно
    прерванный прогон валил отчёт целиком, и читалось это как «сломан отчёт»,
    хотя сломан слой, а снапшоты целы.
    """
    try:
        nights = ledger.read_nights(ledger_dir)
        pace_rows = ledger.read_pace(ledger_dir)
        index = ledger.load_index(ledger_dir) or {}
        capacity = {(item["username"], item["unit"]): item["capacity"]
                    for item in index.get("capacity", [])}
    except ledger.LedgerError as e:
        return [], [], {}, None, str(e)
    except (OSError, EOFError, zlib.error, ValueError, KeyError,
            TypeError) as e:
        return [], [], {}, None, (f"{type(e).__name__}: {e} — файл слоя "
                                  f"повреждён (обычно оборванная дозапись); "
                                  f"снапшоты целы, лечится ledger.rebuild")
    return nights, pace_rows, capacity, index or None, ""


def _fact_months(nights: list[dict], limit: int) -> list[str]:
    months = sorted({row["night"][:7] for row in nights})
    return months[-limit:] if limit else months


def _unit_rows(objects: dict, by_user_unit: dict, capacity: dict,
               ahead_months: list[str], fact_months: list[str],
               regions: dict) -> list[dict]:
    """Строка = юнит. Порядок: как в реестре целей, юниты внутри — по имени."""
    order = {username: i for i, username in enumerate(objects)}
    keys = set(by_user_unit)
    for username, obj in objects.items():
        for unit in (obj.get("units") or {}):
            keys.add((username, unit))

    rows = []
    for username, unit in sorted(
            keys, key=lambda k: (order.get(k[0], len(order)), k[0], k[1])):
        obj = objects.get(username) or {}
        cells = (obj.get("units") or {}).get(unit)
        nights = by_user_unit.get((username, unit), [])
        engine = obj.get("engine") or (nights[0].get("engine") if nights else "")
        granularity = (obj.get("granularity")
                       or (nights[0].get("granularity") if nights else ""))
        cap = capacity.get((username, unit))
        if cap is None and cells:
            cap = core.unit_capacity(cells)
        kind = row_kind(engine, granularity, unit, cap)
        sold, sold_night, no_fund = _nearest_sold(cells, cap)
        row = {
            "username": username,
            "unit": unit,
            "site": obj.get("site") or "",
            "engine": engine or "",
            "granularity": granularity or "",
            "capacity": cap,
            "kind": kind,
            "label": row_label(kind, unit, capacity=cap, sold=sold,
                               night=sold_night, no_fund=no_fund),
            "region": regions.get(username),
            "in_grid": bool(cells),
            "quota": bool(obj.get("quota")
                          or (nights and nights[0].get("source_kind")
                              == "aggregator_quota")),
            "checked_at": obj.get("checked_at"),
            "run_id": obj.get("run_id"),
            "status": obj.get("status"),
            "source_label": obj.get("source_label") or "",
            "ahead": _unit_ahead(cells or {}, ahead_months),
            "fact": _unit_fact(nights, fact_months, cap),
            "fact_cuts": ledger.realized(rows=nights)["cuts"] if nights
                         else {c: {"busy": 0, "known": 0, "pct": None}
                               for c in ledger.CUTS},
            "fact_nights": len({r["night"] for r in nights}),
            "inventory_gap": {m: bool((obj.get("inventory_gap") or {}).get(m))
                              for m in ahead_months},
        }
        if not row["in_grid"]:
            row["source_label"] = (
                "юнита нет в сегодняшней сетке: цель выпала из реестра или "
                "движок её больше не отдаёт; цифры ниже — из ряда")
        row["notes"] = _row_notes(row, ahead_months)
        rows.append(row)
    return rows


def _nearest_sold(cells: Optional[dict], capacity: Optional[int]) -> tuple:
    """«Продано N из M» на ближайшую СНЯТУЮ ночь — то, что отвечает движок.

    -> (продано, ночь, фонда на ближние ночи нет). Ночь, на которую движок
    фонд не назвал, пропускается: `core.cell_units` отдал бы по ней «1 домик»,
    и подпись склеила бы это с фондом из максимума за всю историю — «продано 0
    из 3» там, где про фонд этой ночи не сказано ничего, а при `busy` ещё и
    «продано 1 из 3» там, где тип продан целиком.
    """
    if not cells or capacity is None:
        return None, None, False
    measured = False
    for day in sorted(cells):
        cell = cells[day]
        if cell.get("state") not in ("free", "busy"):
            continue
        measured = True
        total, busy = core.cell_units(cell)
        if total < capacity:
            continue
        return busy, day, False
    return None, None, measured


# ---------------------------------------------------------------------------
# Сезонность: факт по месяцам, медиана рынка, темп
# ---------------------------------------------------------------------------

def _seasonality_rows(by_user_unit: dict, fact_months: list[str],
                      units: list[dict]) -> list[dict]:
    """Факт по объекту и месяцу. Только факт: продажи вперёд сюда не ходят.

    Смешанное основание считается ЗДЕСЬ, а не только в поюнитной ячейке:
    лист «Сезонность по месяцам» — та самая таблица, ради которой писался
    тикет 15, и заказчик читает именно её. `inventory_short` — сколько
    юнит-ночей месяца пришли с фондом меньше фонда юнита: в них тип из трёх
    домиков весит один, и процент занижен методически, а не потому что не
    бронируют.
    """
    per_user: dict[str, list] = {}
    for (username, _unit), rows in by_user_unit.items():
        per_user.setdefault(username, []).extend(rows)
    region_of = {u["username"]: u["region"] for u in units}
    capacity_of = {(u["username"], u["unit"]): u["capacity"] for u in units}
    out = []
    for username in sorted(per_user):
        have = {r["night"][:7] for r in per_user[username]}
        for month in fact_months:
            if month not in have:
                continue
            res = ledger.realized(month=month, rows=per_user[username])
            days = _days_in_month(month)
            measured = [r for r in _kept_nights(per_user[username], month)
                        if r.get("state") in ("free", "busy")]
            short = [r for r in measured
                     if capacity_of.get((r["username"], r["unit"]))
                     and (r.get("units_total") or 0)
                     < capacity_of[(r["username"], r["unit"])]]
            out.append({
                "username": username,
                "region": region_of.get(username),
                "month": month,
                "label": _month_label(month),
                "cuts": res["cuts"],
                "nights": res["nights"],
                "days_in_month": days,
                "complete": res["nights"] >= days,
                "unknown": res["unknown"],
                "sales_not_open": res["sales_not_open"],
                "gap_excluded": res["gap_excluded"],
                "quota_excluded": res["quota_excluded"],
                # знаменатель месяца: юниты ЭТОГО месяца и только те, что
                # реально пошли в факт (квота и опоздавшие ночи выброшены)
                "units": len({r["unit"] for r
                              in _kept_nights(per_user[username], month)}),
                # смешанное основание: юнит-ночи, где фонд снят не был
                "inventory_short": len({(r["unit"], r["night"])
                                        for r in short}),
                "measured_unit_nights": len({(r["unit"], r["night"])
                                             for r in measured}),
                "units_short": len({r["unit"] for r in short}),
            })
    return out


def _market_rows(nights: list[dict], fact_months: list[str], targets,
                 coverage: dict, seasonality: Optional[list] = None
                 ) -> list[dict]:
    """Медиана занятости по объектам. Регионы — НИКОГДА не в одной медиане.

    Медиана складывает проценты объектов, и если у части из них знаменатель
    смешан из двух оснований (часть ночей снята с фондом, часть без него),
    то занижены и они, и медиана. Считается это не заново, а по тем же
    строкам сезонности, что печатаются выше на листе: второй копии признака
    в отчёте быть не должно.
    """
    regions = sorted(coverage["regions"]) if coverage["ok"] else [None]
    out = []
    rows = _latest_nights(nights)
    season = seasonality or []
    for region in regions:
        for month in fact_months:
            res = ledger.market_median(region=region, month=month, rows=rows,
                                       targets=targets)
            # Ночи и полнота — из ТОГО ЖЕ отбора, что и медиана: у самарского
            # ряда в три ночи полноту подписывал подмосковный, и «0.0%»
            # выходило как месячный факт. Плюс сюда не идут ночи, снятые с
            # опозданием, и квоты — их медиана тоже не считает.
            cov = ledger.realized(region=region, month=month, rows=rows,
                                  targets=targets)
            days = _days_in_month(month)
            in_cut = [s for s in season if s["month"] == month
                      and (region is None or s["region"] == region)]
            out.append({
                "region": region or "все цели",
                "mixed_regions": region is None,
                "objects_short": sum(1 for s in in_cut
                                     if s.get("inventory_short")),
                "objects_counted": len(in_cut),
                "month": month,
                "label": _month_label(month),
                "median_pct": res["median_pct"],
                "objects": res["objects"],
                "per_object": res["per_object"],
                "nights": cov["nights"],
                "gap_excluded": cov["gap_excluded"],
                "quota_excluded": cov["quota_excluded"],
                "days_in_month": days,
                "complete": cov["nights"] >= days,
            })
    return out


def _pace_curves(pace_rows: list[dict], units: list[dict],
               months: list[str]) -> list[dict]:
    """Кривая темпа по объекту и месяцу: за сколько суток выкупают ночь."""
    if not pace_rows:
        return []
    wanted = set(months)
    # Точки раскладываются по кривым ОДИН раз. Раньше каждая кривая заново
    # сворачивала весь список: 70 кривых × 39878 точек — 2.79 млн проходов, и
    # обе величины растут вместе (целей станет вдвое больше, горизонт —
    # скользящий год). Тот же приём, которым тикет 09 лечил load_history.
    grouped: dict[tuple, list] = {}
    for row in pace_rows:
        month = row["night"][:7]
        if month in wanted:
            grouped.setdefault((row["username"], month), []).append(row)
    region_of = {u["username"]: u["region"] for u in units}
    out = []
    for username, month in sorted(grouped):
        points = ledger.pace_curve(username, month,
                                   rows=grouped[(username, month)])
        if not points:
            continue
        out.append({"username": username, "region": region_of.get(username),
                    "month": month, "label": _month_label(month),
                    "points": points})
    return out


def night_pace(data: dict, username: str, night: str) -> list[dict]:
    """Как заполнялась ОДНА ночь: точки по lead_days (ручная сверка кривой)."""
    return ledger.pace_series(username, night, rows=data["_pace_rows"])


# ---------------------------------------------------------------------------
# Таблицы
# ---------------------------------------------------------------------------

def unit_headers(data: dict) -> list[str]:
    """Шапка поюнитной таблицы. Одна на все листы — иначе их не сравнить."""
    head = [COL_NUM, COL_OBJECT, COL_SITE, COL_KIND, COL_ENGINE, COL_CAPACITY]
    head += [f"{FACT_PREFIX}{_month_label(m)}" for m in data["months_fact"]]
    if not data["months_fact"]:
        head.append(f"{FACT_PREFIX}ряда ещё нет")
    # Будни/выходные считаны по ВСЕМУ ряду сразу, а не по месяцу: об этом
    # сказано в шапке, иначе читатель сложит их с месячными колонками.
    head += [f"{FACT_PREFIX}будни (весь ряд)", f"{FACT_PREFIX}выходные "
             f"(весь ряд)", COL_NIGHTS]
    head += [f"{AHEAD_PREFIX}{_month_label(m)}" for m in data["months_ahead"]]
    head += [COL_NOTES, COL_CHECKED, COL_STATE, COL_VERDICT, COL_COMMENT]
    return head


def unit_row_cells(data: dict, row: dict, number: int) -> list:
    cells = [number, row["username"], row["site"], row["label"], row["engine"],
             row["capacity"] if row["capacity"] is not None else "не снят"]
    for month in data["months_fact"]:
        cells.append(fact_cell(row["fact"].get(month, {})))
    if not data["months_fact"]:
        cells.append("слой сезонности пуст")
    cells.append(_fmt_pct(row["fact_cuts"]["weekday"]["pct"]))
    cells.append(_fmt_pct(row["fact_cuts"]["weekend"]["pct"]))
    cells.append(row["fact_nights"])
    for month in data["months_ahead"]:
        cells.append(ahead_cell(row["ahead"].get(month, {}),
                                gap=row["inventory_gap"].get(month, False)))
    cells += [row["notes"], _fmt_day(row["checked_at"]), row["source_label"],
              "", ""]
    return cells


def unit_sheet(data: dict, rows: Optional[list] = None) -> tuple:
    """(шапка, строки) для всех юнитов или для переданного подмножества."""
    rows = data["units"] if rows is None else rows
    headers = unit_headers(data)
    return headers, [unit_row_cells(data, row, i)
                     for i, row in enumerate(rows, 1)]


def _unit_groups(data: dict) -> list[tuple]:
    """(имя листа, строки модели) поюнитной таблицы — без рендера ячеек.

    Отдельно от `unit_sheets`, потому что перечень листов нужен и тексту «Как
    читать»: считать ради одного списка имён все ячейки книги незачем.
    """
    if not data["region"]["ok"]:
        return [(ALL_TARGETS_SHEET, data["units"])]
    grouped: dict[Optional[str], list] = {}
    for row in data["units"]:
        grouped.setdefault(row["region"], []).append(row)
    groups = [(f"Юниты — {region}", grouped[region])
              for region in sorted(r for r in grouped if r)]
    if grouped.get(None):
        groups.append((NO_REGION_SHEET, grouped[None]))
    return groups


def unit_sheets(data: dict) -> list[tuple]:
    """Листы поюнитной таблицы: по региону, если разрез доступен.

    Самара и Подмосковье — отдельные листы и отдельные медианы, никогда не
    в одной: это разные рынки, и средняя по ним не значит ничего.
    """
    headers = unit_headers(data)
    return [(name, headers, unit_sheet(data, rows)[1])
            for name, rows in _unit_groups(data)]


SEASON_HEADERS = ["Объект", "Регион", "Месяц", "Занято, % (все ночи)",
                  "будни", "выходные", "Домико-ночей занято",
                  "Домико-ночей известно", "Юнитов", "Ночей в ряду",
                  "Ночей в месяце", "Оговорка", COL_VERDICT, COL_COMMENT]


def _mixed_basis_note(row: dict) -> str:
    """Пометка смешанного знаменателя строки сезонности; чисто — пусто.

    Один текст на лист, выжимку и медиану: пометка, написанная в трёх местах
    по-разному, читается как три разные оговорки.
    """
    short = row.get("inventory_short") or 0
    if not short:
        return ""
    base = row.get("measured_unit_nights") or short
    return (f"фонд снят не на все ночи ({short} из {base} снятых "
            f"юнит-ночей): в них юнит весит меньше своих домиков — процент "
            f"месяца занижен методически, а не потому что не бронируют")


def season_sheet(data: dict) -> tuple:
    rows = []
    for r in data["seasonality"]:
        cuts = r["cuts"]
        note = _mixed_basis_note(r)
        if not r["complete"]:
            note += (f"{'; ' if note else ''}неполный ряд: {r['nights']} "
                     f"{_nights_word(r['nights'])} из {r['days_in_month']} — "
                     f"это не месячный итог")
        if r["gap_excluded"]:
            n = r["gap_excluded"]
            note += (f"{'; ' if note else ''}{n} {_nights_word(n)} "
                     f"{_agree(n, 'снята', 'сняты')} с опозданием и в факт "
                     f"{_agree(n, 'не пошла', 'не пошли')}")
        if r["sales_not_open"]:
            n = r["sales_not_open"]
            note += (f"{'; ' if note else ''}{n} {_nights_word(n)} "
                     f"{_agree(n, 'закрыта', 'закрыты')} продажами")
        if r["quota_excluded"]:
            note += (f"{'; ' if note else ''}объект снят квотой "
                     f"площадки-агрегатора — это доля площадки, а не "
                     f"загрузка объекта, и в факт она не идёт")
        rows.append([r["username"], r["region"] or "не указан", r["label"],
                     _fmt_pct(cuts["all"]["pct"]),
                     _fmt_pct(cuts["weekday"]["pct"]),
                     _fmt_pct(cuts["weekend"]["pct"]),
                     cuts["all"]["busy"], cuts["all"]["known"], r["units"],
                     r["nights"], r["days_in_month"], note, "", ""])
    for m in data["market"]:
        if m.get("mixed_regions"):
            note = ("медиана по ВСЕМ целям сразу: разрез по региону "
                    "недоступен (поле region проставлено не у всех целей), "
                    "поэтому рынки в ней смешаны — Самара и Подмосковье "
                    "попали в одну цифру, и рынком она не является")
        else:
            note = ("медиана по объектам региона, а не средняя по "
                    "домико-ночам: средняя тонет в самом крупном объекте")
        note += ("; в колонке «Юнитов» у этой строки стоит число объектов, "
                 "попавших в медиану")
        if m.get("objects_short"):
            note += (f"; у {m['objects_short']} из {m['objects_counted']} "
                     f"объектов разреза фонд снят не на все ночи — их "
                     f"проценты занижены методически, и медиана вместе с ними")
        if m["objects"] == 0:
            note += "; ни одного объекта с цифрой в этом разрезе"
        if not m["complete"]:
            note += (f"; неполный ряд: {m['nights']} "
                     f"{_nights_word(m['nights'])} из {m['days_in_month']} — "
                     f"это не месячный итог")
        if m.get("gap_excluded"):
            n = m["gap_excluded"]
            note += (f"; {n} {_nights_word(n)} {_agree(n, 'снята', 'сняты')} "
                     f"с опозданием и в медиану {_agree(n, 'не пошла', 'не пошли')}")
        rows.append([f"МЕДИАНА — {m['region']}", m["region"], m["label"],
                     _fmt_pct(m["median_pct"]), "—", "—", "—", "—",
                     m["objects"], m["nights"], m["days_in_month"], note,
                     "", ""])
    return SEASON_HEADERS, rows


def pace_headers() -> list[str]:
    head = ["Объект", "Регион", "Месяц"]
    head += [f"за {lead} сут." if lead else "утром ночи"
             for lead in ledger.LEAD_STEPS]
    head += ["Оговорка", COL_VERDICT, COL_COMMENT]
    return head


def pace_sheet(data: dict) -> tuple:
    headers = pace_headers()
    rows = []
    for r in data["pace"]:
        points = {p["lead_days"]: p for p in r["points"]}
        cells = [r["username"], r["region"] or "не указан", r["label"]]
        for lead in ledger.LEAD_STEPS:
            point = points.get(lead)
            cells.append(pace_cell(point) if point else "")
        weak = [p["lead_days"] for p in r["points"] if p["pct"] is None]
        partial = [p for p in r["points"] if p["pct"] is not None
                   and (p["partial"] or p["inventory_missing"])]
        bits = []
        if weak:
            bits.append("у шагов без цифры причина написана прямо в ячейке "
                        "(мало ночей, ночи не сняты, продажи не открыты, "
                        "фонд не снят) — сравнивать такой шаг с соседними "
                        "нельзя")
        if partial:
            bits.append(f"у {len(partial)} шагов с цифрой набор ночей короче "
                        f"опорного (он назван в скобках): их падение слева "
                        f"направо частью методическое, а не про продажи — "
                        f"как единый ряд читаются только шаги без скобок")
        cells += ["; ".join(bits), "", ""]
        rows.append(cells)
    return headers, rows


LIMITS_HEADERS = ["Чего нет", "Почему", "Когда будет"]


def _timeline(data: dict) -> dict:
    """Сроки отчёта — из ряда на диске, а не из памяти автора.

    Всё, что лист «Чего пока нет» и абзац «Честные сроки» обещают о будущем,
    считается здесь. Статический текст в разделе, который называется
    «честные сроки», устаревает молча и ровно в тот день, когда данные
    появляются: файл начал бы печатать «Самара — ни одной цифры» рядом с
    таблицей, где самарские цифры уже стоят.
    """
    complete = sorted({r["month"] for r in data["seasonality"]
                       if r["complete"]})
    span = data.get("ledger_span") or {}
    first_full = complete[0] if complete else _first_full_month(span.get("first"))
    second = (complete[1] if len(complete) > 1
              else (_month_shift(first_full, 1) if first_full else None))
    return {
        "complete_months": complete,
        "first_night": span.get("first"),
        "last_night": span.get("last"),
        "first_full_month": first_full,
        "second_full_month": second,
        # год — это двенадцать полных месяцев подряд, считая от первого
        "year_month": _month_shift(first_full, 11) if first_full else None,
    }


def _coverage_facts(data: dict) -> dict:
    """Регион реестра -> сколько его целей вообще без сетки и без факта.

    Ключ «» — цели, которым регион ещё не проставлен: они тоже обязаны
    попасть в «чего нет», иначе разрез молча съест их вместе с причиной.
    """
    with_units = {u["username"] for u in data["units"]}
    with_fact = {r["username"] for r in data["seasonality"]
                 if r["cuts"]["all"]["known"]}
    per_region: dict = {}
    for target in data["targets"]:
        if not isinstance(target, dict) or not target.get("username"):
            continue
        name = target["username"]
        stat = per_region.setdefault(target.get("region") or "",
                                     {"targets": 0, "units": 0, "fact": 0,
                                      "no_units": []})
        stat["targets"] += 1
        stat["units"] += name in with_units
        stat["fact"] += name in with_fact
        if name not in with_units:
            stat["no_units"].append(name)
    return per_region


def _region_title(region: str) -> str:
    return f"регион «{region}»" if region else "цели без региона в реестре"


def _month_rows(data: dict) -> list[list]:
    """Строки «чего нет» про месяцы — по тому, что реально лежит в ряду."""
    tl = _timeline(data)
    done = tl["complete_months"]
    rows = []
    if not done:
        if tl["first_full_month"]:
            first = tl["first_night"]
            rows.append([
                "Полный месяц факта",
                f"слежка начата {_fmt_day(first)}, поэтому "
                f"{_month_words(first[:7])} застан не целиком: месячным "
                f"итогом такая цифра не является",
                f"первый полный месяц — "
                f"{_month_words(tl['first_full_month'])}, готов "
                f"{_ready_day(tl['first_full_month'])}"])
        else:
            rows.append(["Полный месяц факта",
                         "прошедших ночей в ряду нет вовсе: слой сезонности "
                         "пуст, факт считать не из чего",
                         "после пересборки слоя (ledger.rebuild) и первого "
                         "прогона"])
    if len(done) < 2 and tl["second_full_month"]:
        rows.append([
            "Сравнение двух месяцев подряд",
            f"полных месяцев в ряду {len(done)}, а сезонность — это разница "
            f"между месяцами: одного для неё мало",
            f"второй полный месяц — {_month_words(tl['second_full_month'])}, "
            f"готов {_ready_day(tl['second_full_month'])}"])
    if len(done) < 12 and tl["year_month"]:
        rows.append([
            "Годовой цикл",
            f"полных месяцев в ряду {len(done)} из 12: полный круг сезонов "
            f"снимается ровно год, короче не бывает",
            f"двенадцатый полный месяц — {_month_words(tl['year_month'])}, "
            f"готов {_ready_day(tl['year_month'])}"])
    return rows


def limits_rows(data: dict) -> list[list]:
    """Решение Р6: честные сроки прямым текстом — и посчитанные по данным."""
    rows = []
    per_region = _coverage_facts(data)
    for region in sorted(per_region):
        stat = per_region[region]
        title = _region_title(region)
        shown = ", ".join(stat["no_units"][:5])
        if not stat["units"]:
            rows.append([
                f"{title} — ни одной цифры",
                f"целей в реестре {stat['targets']}, сетка не снята ни у "
                f"одной: рецепт съёма не разведан или модуль не отвечает "
                f"({shown})",
                "первый снимок — сразу после разведки рецептов"])
        elif not stat["fact"]:
            rows.append([
                f"{title} — фактическая занятость",
                f"сетка снята у {stat['units']} целей из {stat['targets']}, "
                f"но прошедших ночей под слежкой ещё нет: факт появляется "
                f"только у ночи, которую слежка застала",
                "первая цифра факта — после первой же прошедшей ночи"])
    rows += _month_rows(data)
    if not data["region"]["ok"]:
        rows.insert(0, ["Разрез «Самара отдельно от Подмосковья»",
                        data["region"]["note"],
                        "как только поле region проставлено в реестре целей"])
    if data.get("ledger_hole"):
        # Дыра в слое — это дыра в цифрах ЭТОГО файла: ряд без пропущенных
        # суток выглядит как честный, и о пропуске нельзя узнать ниоткуда,
        # кроме index.json. Формулировку даёт сам слой (ledger.hole_note),
        # чтобы отчёт и лента прогона не расходились словами.
        rows.insert(0, ["Ночи прогонов, не попавших в ряд",
                        f"слой сезонности неполон: {data['ledger_hole']}",
                        "сразу после пересборки слоя (ledger.rebuild): "
                        "снапшоты целы, цифры вернутся"])
    if data["ledger_error"]:
        rows.insert(0, ["Фактическая занятость по всем объектам",
                        f"слой сезонности не прочитан: {data['ledger_error']}",
                        "после пересборки слоя (ledger.rebuild)"])
    return rows


# ---------------------------------------------------------------------------
# Лист «Как читать»
# ---------------------------------------------------------------------------

def how_to_read_lines(data: dict) -> list[str]:
    """Человеческий текст: зачем это, что здесь и чему верить нельзя."""
    day = data["generated_for"].strftime("%d.%m.%Y")
    by_kind = data["by_kind"]
    fact = data["months_fact"]
    lines = [
        f"Глэмпинги по юнитам — снятие {day}. Одна строка = один юнит модуля "
        f"бронирования, а не объект целиком.",
        "",
        "ЗАЧЕМ. Ты трижды просил показать, какие домики сдаются, а какие нет: "
        "месячный процент по объекту слишком размытый, по нему не видно, кто "
        "из домиков простаивает. Здесь каждая строка — отдельный юнит, и по "
        "ней сразу видно и прошедшие ночи, и то, что уже продано вперёд.",
        "",
        "ТРИ РОДА СТРОК, И ЭТО ВАЖНО. Модули бронирования устроены "
        "по-разному, поэтому строки честно называются по-разному, и путать их "
        "нельзя. «Домик: <имя>» — движок адресует конкретный физический "
        "номер, и строка правда про один домик. «Тип: <имя> — продано N из "
        "M» — движок торгует категорией: за именем стоит M одинаковых "
        "домиков, и он отвечает только «сколько из них ещё свободно». «Тип: "
        "<имя> — фонд не снят» — сколько домиков за именем, движок не "
        "говорит вовсе, и такая строка занята только тогда, когда продан "
        "последний домик; её процент занижен.",
        "",
        f"Сейчас в таблице: домиков {by_kind[KIND_HOME]}, типов с известным "
        f"фондом {by_kind[KIND_TYPE]}, типов без снятого фонда "
        f"{by_kind[KIND_TYPE_NO_CAPACITY]}. Род определяется движком и тем, "
        f"снят ли фонд, а НЕ именем: юнит может называться «Сфера 1» и быть "
        f"типом из пяти сфер, а может называться «Этнодом» и быть четырьмя "
        f"домиками сразу. Поэтому строка второго рода не называется домиком "
        f"даже тогда, когда домик в ней ровно один.",
        "",
        "ДВЕ РАЗНЫЕ ЦИФРЫ, КОТОРЫЕ НЕЛЬЗЯ СКЛАДЫВАТЬ. Колонки «Факт: <месяц>» "
        "— это реализованная занятость уже прошедших ночей: чем ночь "
        "кончилась, снято утром самой ночи и больше не меняется. Колонки "
        "«Продано вперёд: <месяц>» — это глубина продаж на будущие ночи по "
        "сегодняшнему снимку: сколько уже выкуплено из того, что ещё "
        "впереди. Первая говорит, как объект отработал; вторая — как он "
        "продаётся сейчас. В одну колонку они не сводятся никогда.",
        "",
        "ЧЕМУ ВЕРИТЬ НЕЛЬЗЯ, И ЭТО НАПИСАНО В САМОЙ ЯЧЕЙКЕ. Если ряд короче "
        "месяца, в ячейке стоит «неполный ряд» и сказано, по скольким ночам "
        "из скольких посчитано: слежка идёт не с первого числа, и месячным "
        "итогом такая цифра не является. Если у месяца снят фонд не на все "
        "ночи, об этом сказано прямо в ячейке: на дальних ночах тип из трёх "
        "домиков весит один, и процент занижен методически, а не потому что "
        "не бронируют. Если модуль ещё не открыл продажи на месяц, стоит "
        "«продажи не открыты» — это пустой календарь, а не аншлаг.",
        "",
        _sheets_sentence(data),
        "",
        _deadlines_sentence(data),
        "",
        "ТВОИ КОЛОНКИ — «Вердикт заказчика» (следить / вести / пропустить / "
        "вопрос) и «Комментарий». Они пустые, заполнять их никто, кроме "
        "тебя, не будет.",
    ]
    if fact:
        lines += ["", f"Ряд фактических ночей сейчас накоплен за "
                      f"{', '.join(_month_label(m) for m in fact)}; "
                      f"продажи вперёд показаны за "
                      f"{', '.join(_month_label(m) for m in data['months_ahead'])}."]
    else:
        lines += ["", "Слой сезонности пока пуст, поэтому колонок факта в "
                      "таблице нет — есть только глубина продаж вперёд. "
                      "Лечится пересборкой слоя, данные для него уже сняты."]
    if data.get("ledger_hole"):
        lines += ["", "ВНИМАНИЕ: в ряду не хватает суток, и цифры факта ниже "
                      "посчитаны без них. Что именно: " + data["ledger_hole"]
                      + ". Снапшоты целы, лечится пересборкой слоя."]
    if data["ledger_error"]:
        lines += ["", "ВНИМАНИЕ: слой сезонности не прочитался, и весь факт в "
                      "этом файле отсутствует. Причина: "
                      + data["ledger_error"]
                      + ". Пересобрать слой — и факт вернётся, снапшоты целы."]
    if not data["region"]["ok"]:
        lines += ["", "РАЗРЕЗ «САМАРА ОТДЕЛЬНО ОТ ПОДМОСКОВЬЯ» ПОКА НЕ "
                      "СЧИТАЕТСЯ, и вот почему: " + data["region"]["note"]]
    return lines


def _deadlines_sentence(data: dict) -> str:
    """Абзац «честные сроки» — из ряда на диске, а не из памяти автора.

    Раздел, который называется честным, обязан быть посчитан: обещание
    «первый честный месяц — октябрь» в декабре обесценивает весь механизм
    оговорок, на котором держится отчёт.
    """
    tl = _timeline(data)
    done = tl["complete_months"]
    bits = ["ЧЕСТНЫЕ СРОКИ."]
    gaps = []
    for region, stat in sorted(_coverage_facts(data).items()):
        if not stat["units"]:
            gaps.append(f"{_region_title(region)} — сетка не снята ни у одной "
                        f"из {stat['targets']} целей, рецепты не разведаны")
        elif not stat["fact"]:
            gaps.append(f"{_region_title(region)} — сетка снята, но ни одной "
                        f"прошедшей ночи под слежкой ещё нет")
    if gaps:
        bits.append("Цифр сегодня нет вот по кому: " + "; ".join(gaps) + ".")
    if not done and tl["first_full_month"]:
        bits.append(f"Полных месяцев факта в ряду пока нет: слежка начата "
                    f"{_fmt_day(tl['first_night'])}, поэтому "
                    f"{_month_words(tl['first_night'][:7])} обрезан её "
                    f"началом и месячным итогом быть не может. Первый полный "
                    f"месяц — {_month_words(tl['first_full_month'])}, он "
                    f"будет готов {_ready_day(tl['first_full_month'])}.")
    elif done:
        bits.append("Полных месяцев факта в ряду "
                    + str(len(done)) + ": "
                    + ", ".join(_month_words(m) for m in done)
                    + " — только их и можно читать как месячный итог.")
    else:
        bits.append("Прошедших ночей в ряду нет вовсе: слой сезонности пуст, "
                    "и факта в этом файле не будет, пока слой не пересобран.")
    if len(done) < 2 and tl["second_full_month"]:
        bits.append(f"Сравнить два месяца подряд — то есть увидеть саму "
                    f"сезонность — можно будет "
                    f"{_ready_day(tl['second_full_month'])}.")
    if len(done) < 12 and tl["year_month"]:
        bits.append(f"Полный годовой цикл соберётся "
                    f"{_ready_day(tl['year_month'])}: короче года сезонность "
                    f"не снимается ничем.")
    return " ".join(bits)


def _sheet_names(data: dict) -> list[str]:
    """Имена листов собранного файла — ровно те, что окажутся в xlsx.

    Считаются тем же кодом, что и пишет книгу: обещанный в тексте лист и
    настоящий лист обязаны совпадать даже после обрезки и разведения имён.
    """
    names = [HOW_SHEET.split(" ", 1)[-1]]
    names += [name for name, _rows in _unit_groups(data)]
    if data["no_data"]:
        names.append(NODATA_SHEET)
    names += [SEASON_SHEET, PACE_SHEET, LIMITS_SHEET]
    return _sheet_titles(names)


def _sheets_sentence(data: dict) -> str:
    """Перечень листов — по факту собранного файла, а не по памяти автора."""
    bits = ["Поюнитные таблицы (по одной на регион, если регион проставлен)"]
    if data["no_data"]:
        bits.append(f"«{NODATA_SHEET}» — цели, по которым цифр нет, с "
                    f"причиной у каждой (сейчас их {len(data['no_data'])})")
    bits.append(f"«{SEASON_SHEET}» — факт по объектам и медиана по региону")
    bits.append(f"«{PACE_SHEET}» — за сколько суток до ночи её выкупают "
                f"(это и есть окно, в котором работает маркетинг)")
    bits.append(f"«{LIMITS_SHEET}» — что мы честно не можем показать и "
                f"когда сможем")
    return "ЛИСТЫ. " + "; ".join(bits) + "."


# ---------------------------------------------------------------------------
# Рендеры
# ---------------------------------------------------------------------------

def _md_table(headers: list[str], rows: list[list]) -> list[str]:
    out = ["| " + " | ".join(md_cell(h) for h in headers) + " |",
           "|" + "|".join("---" for _ in headers) + "|"]
    for row in rows:
        out.append("| " + " | ".join(md_cell(c) for c in row) + " |")
    return out


def render_markdown(data: dict) -> str:
    """Версионируемая выжимка рядом с исследованием рынка глэмпингов."""
    day = data["generated_for"].strftime("%d.%m.%Y")
    by_kind = data["by_kind"]
    out = [f"# Глэмпинги по юнитам и сезонность — снятие {day}", "",
           "Отчёт на прямую просьбу заказчика (26.08, повторена 27.08 и 28.08): "
           "«мы не видим, какие домики сдаются, какие нет; месячные проценты "
           "слишком размыты». Одна строка = один юнит модуля бронирования.",
           "",
           "Файлы этой редакции:",
           f"- xlsx для заказчика — `agent-runtime/research/glamping/occupancy/"
           f"out/глэмпинги-по-юнитам-{data['generated_for'].isoformat()}"
           f".xlsx`; листы: "
           + ", ".join(f"«{n}»" for n in _sheet_names(data)) + ";",
           "- этот файл — версионируемая выжимка; сводка «снимок перед "
           "интервью» живёт отдельно в `occupancy.md` и не трогается.",
           "",
           "## Три рода строк", "",
           "Род определяется движком и наличием снятого фонда, а не именем "
           "юнита: «Сфера 1» может быть типом из пяти сфер, а «Этнодом» — "
           "четырьмя домиками под одним именем. Строка второго рода не "
           "называется домиком даже при одном домике.", ""]
    out += _md_table(["Род строки", "Подпись в таблице", "Строк"],
                     [[KIND_TITLES[KIND_HOME], "Домик: <имя>",
                       by_kind[KIND_HOME]],
                      [KIND_TITLES[KIND_TYPE], "Тип: <имя> — продано N из M",
                       by_kind[KIND_TYPE]],
                      [KIND_TITLES[KIND_TYPE_NO_CAPACITY],
                       "Тип: <имя> — фонд не снят",
                       by_kind[KIND_TYPE_NO_CAPACITY]]])
    total_units = len(data["units"])
    homes = sum(u["capacity"] or 1 for u in data["units"])
    out += ["", f"Всего строк: {total_units}; продаваемых позиций за ними — "
                f"{homes} (юнит без снятого фонда считается за одну, то есть "
                f"это нижняя оценка). Домиками они названы условно: движок "
                f"отдаёт всё, что продаёт, включая нежилое — стоянку для "
                f"катера, баню; отдельного рода для «не жильё» у отчёта нет. "
                f"Объектов: {len({u['username'] for u in data['units']})}.",
            ""]

    out += ["## Факт и продажи вперёд — разные цифры", "",
            "«Фактически занято» — реализованная занятость прошедших ночей из "
            "ряда: состояние ночи, снятое утром самой ночи, дальше оно не "
            "меняется. «Продано вперёд» — глубина продаж на будущие ночи из "
            "последнего снимка. В одну колонку они не сводятся.", ""]
    if data["months_fact"]:
        out += _md_table(["Месяц", "Регион", "Медиана по объектам", "Объектов",
                          "Ночей в ряду", "Полный месяц?", "Основание"],
                         [[m["label"], m["region"], _fmt_pct(m["median_pct"]),
                           m["objects"], m["nights"],
                           "да" if m["complete"] else "нет — неполный ряд",
                           (f"у {m['objects_short']} из "
                            f"{m['objects_counted']} объектов фонд снят не "
                            f"на все ночи — медиана занижена вместе с ними")
                           if m.get("objects_short") else "фонд снят"]
                          for m in data["market"]])
    else:
        out += ["Ряда фактических ночей пока нет: слой сезонности пуст."]
    out += [""]

    out += ["## Сезонность по объектам", ""]
    if data["seasonality"]:
        # Колонка региона здесь не украшение: без неё самарские и
        # подмосковные объекты встают одним списком, а тикет 14 это прямо
        # запрещает. Оговорка — та же, что на листе xlsx.
        out += _md_table(["Объект", "Регион", "Месяц", "Занято, %", "будни",
                          "выходные", "Ночей в ряду", "Полный месяц?",
                          "Оговорка"],
                         [[r["username"], r["region"] or "не указан",
                           r["label"],
                           _fmt_pct(r["cuts"]["all"]["pct"]),
                           _fmt_pct(r["cuts"]["weekday"]["pct"]),
                           _fmt_pct(r["cuts"]["weekend"]["pct"]),
                           f"{r['nights']} из {r['days_in_month']}",
                           "да" if r["complete"] else "нет",
                           _mixed_basis_note(r) or "—"]
                          for r in data["seasonality"]])
    else:
        out += ["Ряда пока нет."]
    out += [""]

    out += ["## Темп бронирования", "",
            "За сколько суток до ночи её выкупают. Читать по убыванию lead: "
            "это и есть окно, в котором работает маркетинг. Шаг без цифры "
            "посчитан по слишком малому числу ночей либо без снятого фонда — "
            "сравнивать его с соседними нельзя.", ""]
    if data["pace"]:
        leads = [0, 3, 7, 14, 21, 30]
        out += _md_table(["Объект", "Месяц"]
                         + [f"за {l} сут." if l else "утром ночи"
                            for l in leads],
                         [[r["username"], r["label"]]
                          + [pace_cell(next((p for p in r["points"]
                                             if p["lead_days"] == l), {}))
                             for l in leads]
                          for r in data["pace"]])
    else:
        out += ["Кривой пока нет: слой темпа пуст."]
    out += [""]

    if data["no_data"]:
        out += ["## Объекты, по которым цифр нет", "",
                "Ни одна из этих целей не исчезла из отчёта: у каждой "
                "написано, почему цифр нет.", ""]
        out += _md_table(NODATA_HEADERS[:-2],
                         [row[:-2] for row in nodata_sheet(data)[1]])
        out += [""]

    out += ["## Чего пока нет и когда будет", ""]
    out += _md_table(LIMITS_HEADERS, limits_rows(data))
    out += ["", "## Как читать таблицу", ""]
    out += [line for line in how_to_read_lines(data)]
    return "\n".join(out).rstrip() + "\n"


# ---------------------------------------------------------------------------
# xlsx
# ---------------------------------------------------------------------------

# Excel не принимает эти знаки в имени листа: openpyxl отвечает ValueError
# уже после того, как все цифры посчитаны. Значение region правит человек
# руками, и одна косая черта в реестре роняла бы сборку целиком.
SHEET_BAD_CHARS = set("/\\?*[]:")


def _sheet_title(name: str) -> str:
    """Имя листа, которое Excel точно примет: без запретных знаков и ≤31."""
    clean = "".join(" " if ch in SHEET_BAD_CHARS else ch for ch in name)
    clean = " ".join(clean.split()).strip("'")[:31].strip()
    return clean or "Лист"


def _sheet_titles(names: list[str]) -> list[str]:
    """Те же имена, но заведомо различные: два длинных региона совпадают в 31 знаке."""
    out, used = [], set()
    for name in names:
        title = _sheet_title(name)
        if title in used:
            for i in range(2, 100):
                tail = f" ({i})"
                candidate = title[:31 - len(tail)].strip() + tail
                if candidate not in used:
                    title = candidate
                    break
        used.add(title)
        out.append(title)
    return out


def write_xlsx(data: dict, path) -> Path:
    """Собрать xlsx в формате, который заказчик принял 04.09 по хаусботам."""
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter

    head_fill = PatternFill("solid", fgColor="1F3864")
    head_font = Font(bold=True, color="FFFFFF", size=10)

    def table(ws, headers, rows, widths=None):
        ws.append(list(headers))
        for row in rows:
            ws.append(list(row))
        for i, name in enumerate(headers, 1):
            cell = ws.cell(row=1, column=i)
            cell.fill, cell.font = head_fill, head_font
            cell.alignment = Alignment(wrap_text=True, vertical="center")
            ws.column_dimensions[get_column_letter(i)].width = (
                widths[i - 1] if widths else _column_width(name))
        ws.freeze_panes = "D2"
        ws.auto_filter.ref = (f"A1:{get_column_letter(len(headers))}"
                              f"{max(len(rows), 1) + 1}")
        for row in ws.iter_rows(min_row=2, max_row=len(rows) + 1):
            for cell in row:
                cell.alignment = Alignment(wrap_text=True, vertical="top")

    wb = openpyxl.Workbook()
    sheets = _unit_groups(data)
    titles = _sheet_titles(
        [HOW_SHEET] + [name for name, _rows in sheets]
        + ([NODATA_SHEET] if data["no_data"] else [])
        + [SEASON_SHEET, PACE_SHEET, LIMITS_SHEET])
    titles.reverse()          # дальше имена разбираются с головы списка
    guide = wb.active
    guide.title = titles.pop()
    guide.column_dimensions["A"].width = 125
    for i, text in enumerate(how_to_read_lines(data), 1):
        cell = guide.cell(row=i, column=1, value=text)
        cell.alignment = Alignment(wrap_text=True, vertical="top")
        if i == 1:
            cell.font = Font(bold=True, size=12)

    headers = unit_headers(data)
    for _name, rows in sheets:
        table(wb.create_sheet(titles.pop()), headers,
              unit_sheet(data, rows)[1])
    if data["no_data"]:
        table(wb.create_sheet(titles.pop()), *nodata_sheet(data))
    table(wb.create_sheet(titles.pop()), *season_sheet(data))
    table(wb.create_sheet(titles.pop()), *pace_sheet(data))
    table(wb.create_sheet(titles.pop()), LIMITS_HEADERS,
          limits_rows(data), widths=[46, 80, 46])

    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(path)
    return path


def _column_width(name: str) -> int:
    if name == COL_NUM:
        return 5
    if name in (COL_NOTES, COL_STATE):
        return 52
    if name in (COL_KIND, COL_SITE):
        return 42
    if name.startswith((FACT_PREFIX, AHEAD_PREFIX)):
        return 30
    if name in (COL_VERDICT, COL_COMMENT):
        return 18
    return 22


# ---------------------------------------------------------------------------
# CLI
# ---------------------------------------------------------------------------

def main(argv=None) -> int:
    parser = argparse.ArgumentParser(
        description="Поюнитный отчёт по загрузке глэмпингов и сезонность")
    parser.add_argument("--snapshot-root", default=None)
    parser.add_argument("--targets", default=None)
    parser.add_argument("--recipes", default=None)
    parser.add_argument("--ledger-dir", default=None)
    parser.add_argument("--out-dir", default=None,
                        help="куда положить xlsx (по умолчанию "
                             "agent-runtime/.../occupancy/out)")
    parser.add_argument("--md", default=None,
                        help="куда положить версионируемую выжимку")
    parser.add_argument("--today", default=None, help="ISO-дата (шов для тестов)")
    parser.add_argument("--months-ahead", type=int, default=3)
    parser.add_argument("--months-fact", type=int, default=DEFAULT_FACT_MONTHS)
    args = parser.parse_args(argv)

    today = date.fromisoformat(args.today) if args.today else core.today()
    data = collect(snapshot_root=args.snapshot_root, targets=args.targets,
                   recipes=args.recipes, ledger_dir=args.ledger_dir,
                   today_=today, months_ahead=args.months_ahead,
                   months_fact=args.months_fact)

    out_dir = Path(args.out_dir or DEFAULT_OUT_DIR)
    xlsx = out_dir / f"глэмпинги-по-юнитам-{today.isoformat()}.xlsx"
    write_xlsx(data, xlsx)
    md_path = Path(args.md or (DEFAULT_MD_DIR / f"units-{today.isoformat()}.md"))
    md_path.parent.mkdir(parents=True, exist_ok=True)
    md_path.write_text(render_markdown(data), encoding="utf-8")

    by_kind = data["by_kind"]
    print(f"OK: {xlsx}")
    print(f"OK: {md_path}")
    print(f"строк: {len(data['units'])} "
          f"(домиков {by_kind[KIND_HOME]}, типов с фондом {by_kind[KIND_TYPE]}, "
          f"типов без фонда {by_kind[KIND_TYPE_NO_CAPACITY]}); "
          f"объектов {len({u['username'] for u in data['units']})}, "
          f"без цифр {len(data['no_data'])}; "
          f"месяцев факта {len(data['months_fact'])}; "
          f"кривых темпа {len(data['pace'])}")
    if data["ledger_hole"]:
        print(f"ВНИМАНИЕ: {data['ledger_hole']}")
    if data["ledger_error"]:
        print(f"ВНИМАНИЕ: слой сезонности не прочитан: {data['ledger_error']}")
    if not data["region"]["ok"]:
        print(data["region"]["note"])
    return 0


if __name__ == "__main__":
    sys.exit(main())
